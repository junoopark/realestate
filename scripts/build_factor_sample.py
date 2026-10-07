"""DFMBA 전처리본(데이터취합_전처리_YYYYMMDD.xlsx)에서 동인 탭용 샘플 데이터를 뽑는다.

원본: https://github.com/ksroh1913/DFMBA 의 impute/build_preprocessed.py 가 만든 엑셀
      (1차_결측보완 시트 = 기준기간 값, 빈칸만 보완)
출력: frontend/data/factor_sample.json

샘플 범위 (화면 확인용으로 일부만 담는다)
  - 시도: 서울·경기·부산  /  서울 자치구: 강남구·마포구·노원구
  - 월별: 2016-01 ~ 최신월,  분기·반기·연간: 2016 ~ 최신
  - 지역에 관계없이 값이 같은 열(금리·통화량 등)은 '전국' 한 계열로 줄인다

사용법:
  python scripts/build_factor_sample.py <데이터취합_전처리_YYYYMMDD.xlsx 경로>
"""

import json
import os
import re
import sys

import openpyxl

SIDO = ["서울", "경기", "부산"]
GU = ["강남구", "마포구", "노원구"]
START_MONTH = 201601
START_YEAR = 2016
PROVISIONAL_VARS = ("V006", "V007")
# 모든 지역에 같은 값이 붙어 있지만 '전국' 값이 아닌 변수: 권역별 조사값을 지역 공통으로 쓴 것
COMMON_LABEL = {"V061": "서울 권역", "V062": "6대광역시 권역", "V063": "기타도시 권역"}

OUT = os.path.join(os.path.dirname(__file__), "..", "frontend", "data", "factor_sample.json")

# 열 이름 끝의 단위 (긴 것부터 맞춘다)
UNITS = ["1980.01.04=100", "2020=100", "천세대당", "전국=1", "십억원", "천명", "천㎡", "천불", "천원",
         "연%", "%", "건", "명", "세대", "호", "원", "동", "㎡", "지수"]


def split_column(col):
    """'V003_아파트_전세가율_%' -> ('V003', '아파트 전세가율', '%')"""
    m = re.match(r"^(?:Y_|서울_)?(V\d{3})_(.*)$", col)
    if not m:
        return None
    vid, rest = m.groups()
    unit = ""
    for u in UNITS:
        if rest.endswith("_" + u):
            unit, rest = u, rest[: -len(u) - 1]
            break
    label = rest.replace("_", " ").strip()
    if unit == "지수":
        unit = ""
    return vid, label, unit


def rnd(v):
    if v is None:
        return None
    return float(f"{v:.6g}")


def sheet_rows(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    header = next(rows)
    return header, [r for r in rows]


def trim(periods, series):
    """앞뒤로 모든 지역이 빈칸인 기간을 잘라낸다"""
    filled = [i for i in range(len(periods)) if any(s[i] is not None for s in series.values())]
    if not filled:
        return periods, series
    a, b = filled[0], filled[-1] + 1
    return periods[a:b], {k: v[a:b] for k, v in series.items()}


def monthly(wb, sheet, regions, scope):
    header, rows = sheet_rows(wb, sheet)
    all_months = sorted({r[3] for r in rows})
    by_region = {reg: {r[3]: r for r in rows if r[1] == reg} for reg in regions}
    metrics = []
    for i, col in enumerate(header):
        if not col or col.startswith(("panel", "region", "month", "D_", "Y_평가")):
            continue
        parts = split_column(col)
        if not parts:
            continue
        vid, label, unit = parts
        # 샘플 기간(START_MONTH~)에 값이 없으면(예: 2015.06 종료된 V034) 전체 기간을 담는다
        months = [m for m in all_months if m >= START_MONTH]
        pick = lambda ms: {reg: [rnd(by_region[reg].get(m, [None] * len(header))[i]) for m in ms] for reg in regions}
        series = pick(months)
        if all(v is None for s in series.values() for v in s):
            months = all_months
            series = pick(months)
        # 서울_ 접두 열은 서울 시 단위 값을 구 시트에 복사해 둔 것이므로 '서울' 한 계열로 둔다
        this_scope = scope
        if col.startswith("서울_"):
            this_scope, series = "seoul", {"서울": series[regions[0]]}
        elif len({json.dumps(s) for s in series.values()}) == 1:
            this_scope, series = "nation", {COMMON_LABEL.get(vid, "전국"): series[regions[0]]}
        if all(v is None for s in series.values() for v in s):
            continue
        periods, series = trim([f"{m // 100}-{m % 100:02d}" for m in months], series)
        metric = dict(var=vid, column=col, label=label, unit=unit, freq="M", scope=this_scope,
                      periods=periods, series=series)
        # 국토부 실거래 신고건수(V006·V007): 원 파이프라인(preprocess/prep_rtms_counts.py) 설명대로
        # 최근 2~3개월은 신고기한(30일)·해제 시차로 값이 계속 바뀐다 -> 마지막 3개월을 잠정으로 표시
        if vid in PROVISIONAL_VARS and len(periods) >= 3:
            metric["provisional_from"] = periods[-3]
        metrics.append(metric)
    return metrics


def lowfreq(wb, sheet, freq):
    """분기_/반기_/연간_ 시트: 시도·구 열로 지역을 구분한다 (전국 행은 '전국')."""
    header, rows = sheet_rows(wb, sheet)
    metrics = []
    for i, col in enumerate(header):
        if i < 3 or not col:
            continue
        parts = split_column(col)
        if not parts:
            continue
        vid, label, unit = parts
        sel = {}
        for r in rows:
            sido, gu, period = r[0], r[1], str(r[2])
            if int(period[:4]) < START_YEAR or r[i] is None:
                continue
            region = gu if gu else sido
            if (gu and gu in GU) or (not gu and (sido in SIDO or sido == "전국")):
                sel.setdefault(region, {})[period] = rnd(r[i])
        if not sel:
            continue
        periods = sorted({p for s in sel.values() for p in s})
        scope = "gu" if any(k in GU for k in sel) else ("nation" if list(sel) == ["전국"] else "sido")
        order = GU if scope == "gu" else ["전국"] + SIDO
        series = {reg: [sel[reg].get(p) for p in periods] for reg in order if reg in sel}
        metrics.append(dict(var=vid, column=col, label=label, unit=unit, freq=freq, scope=scope,
                            periods=periods, series=series))
    return metrics


def notes(wb):
    """2_데이터목록: 열마다 출처·전처리 방법·보완 셀 수"""
    header, rows = sheet_rows(wb, "2_데이터목록")
    idx = {h: i for i, h in enumerate(header)}
    out = {}
    for r in rows:
        col = r[idx["컬럼명"]]
        if col:
            out[col] = dict(source=r[idx["출처(기관·표)"]], method=r[idx["전처리 방법"]],
                            filled=r[idx["보완 셀 수"]], release=r[idx["공표 시기 (공식 근거)"]])
    return out


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    metrics = (monthly(wb, "17시도_1차_결측보완", SIDO, "sido")
               + monthly(wb, "서울25구_1차_결측보완", GU, "gu")
               + lowfreq(wb, "분기_1차_결측보완", "Q")
               + lowfreq(wb, "반기_1차_결측보완", "H")
               + lowfreq(wb, "연간_1차_결측보완", "A"))
    # 같은 열이 시도·구 시트 양쪽에서 '전국'으로 나오면 하나만 남긴다
    seen, unique = set(), []
    for m in metrics:
        key = (m["column"], m["scope"]) if m["scope"] != "nation" else (m["column"], "nation")
        if key in seen:
            continue
        seen.add(key)
        unique.append(m)
    info = notes(wb)
    for m in unique:
        n = info.get(m["column"]) or info.get(m["column"].replace("서울_", "", 1)) or {}
        m.update({k: v for k, v in n.items() if v not in (None, "", "-")})
    overview = dict(wb["1_개요"].iter_rows(values_only=True))
    out = dict(
        _source=("DFMBA 데이터취합_전처리 엑셀(" + os.path.basename(path) + ")의 1차_결측보완 시트. "
                 "샘플: 시도 " + "·".join(SIDO) + ", 서울 구 " + "·".join(GU) + f", {START_YEAR}년~"),
        generated=str(overview.get("생성", "")),
        # 원 파이프라인(variables.py)에서 같은 표로 처리해 별도 열이 없는 변수: V033 = V005 (R-ONE A_2024_00078)
        aliases={"V033": "V005"},
        regions=dict(sido=SIDO, gu=GU),
        metrics=unique,
    )
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{len(unique)}개 계열 -> {os.path.normpath(OUT)} ({os.path.getsize(OUT) / 1024:.0f} KB)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
